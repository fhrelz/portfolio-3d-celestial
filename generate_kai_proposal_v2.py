import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def generate_ubsi_kai_proposal_v2():
    doc = docx.Document()
    
    # Page setup A4 with exact margins: Left 4 cm, Top 3 cm, Right 3 cm, Bottom 3 cm
    section = doc.sections[0]
    section.page_width = Inches(21.0 / 2.54)   # A4 Width 21 cm
    section.page_height = Inches(29.7 / 2.54)  # A4 Height 29.7 cm
    
    section.top_margin = Inches(3.0 / 2.54)
    section.bottom_margin = Inches(3.0 / 2.54)
    section.left_margin = Inches(4.0 / 2.54)
    section.right_margin = Inches(3.0 / 2.54)
    
    # Usable width: 21.0 - 4.0 - 3.0 = 14.0 cm = 7937 dxa
    usable_width_dxa = 7937

    # Helper functions
    def add_body_p(text, indent=True, bold_prefix=None, space_after=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        if indent:
            p.paragraph_format.first_line_indent = Inches(1.25 / 2.54)
            
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = 'Times New Roman'
            r_bold.font.size = Pt(12)
            r_bold.font.bold = True
            
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        return p

    def add_bullet_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(1.25 / 2.54)
        p.paragraph_format.first_line_indent = Inches(-0.6 / 2.54)
        
        r_bullet = p.add_run("•  ")
        r_bullet.font.name = 'Times New Roman'
        r_bullet.font.size = Pt(12)
        
        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.font.name = 'Times New Roman'
            r_b.font.size = Pt(12)
            r_b.font.bold = True
            
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        return p

    def add_bab_header(bab_title, sub_title):
        p1 = doc.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.line_spacing = 1.5
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(bab_title)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(14)
        r1.font.bold = True
        
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.line_spacing = 1.5
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(12)
        r2 = p2.add_run(sub_title)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(14)
        r2.font.bold = True

    def add_subbab_header(title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(title)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.font.bold = True
        return p

    def add_toc_line(label, page_str, is_bold=False, indent_cm=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if indent_cm > 0:
            p.paragraph_format.left_indent = Inches(indent_cm / 2.54)
            
        # Add Tab Stop with Dot Leader aligned at 14.0 cm (7937 dxa)
        pPr = p._p.get_or_add_pPr()
        tabs = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:leader="dot" w:pos="{usable_width_dxa}"/></w:tabs>')
        pPr.append(tabs)
        
        r1 = p.add_run(label)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(12)
        if is_bold:
            r1.font.bold = True
            
        r_tab = p.add_run("\t")
        r_tab.font.name = 'Times New Roman'
        r_tab.font.size = Pt(12)
        
        r2 = p.add_run(page_str)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(12)
        if is_bold:
            r2.font.bold = True

    # ==================== 1. COVER PAGE ====================
    p_cover_title = doc.add_paragraph()
    p_cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cover_title.paragraph_format.line_spacing = 1.5
    p_cover_title.paragraph_format.space_before = Pt(0)
    p_cover_title.paragraph_format.space_after = Pt(0)
    
    r = p_cover_title.add_run("PROPOSAL\nPRAKTIK KERJA LAPANGAN (PKL) / MAGANG\nJUNIOR WEB DEVELOPER & IT SYSTEMS INTERN\nDI PT KERETA API INDONESIA (PERSERO) DAOP 6 YOGYAKARTA")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    
    p_sp1 = doc.add_paragraph()
    p_sp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sp1.paragraph_format.space_before = Pt(18)
    p_sp1.paragraph_format.space_after = Pt(18)
    
    if os.path.exists('bsi_clean.png'):
        r_logo = p_sp1.add_run()
        r_logo.add_picture('bsi_clean.png', width=Inches(2.2))
        
    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author.paragraph_format.line_spacing = 1.5
    p_author.paragraph_format.space_before = Pt(12)
    p_author.paragraph_format.space_after = Pt(24)
    
    r_auth = p_author.add_run("Disusun Oleh:\n\nMUHAMMAD FAHREL YULIYANTO\nNIM : 19240513")
    r_auth.font.name = 'Times New Roman'
    r_auth.font.size = Pt(12)
    r_auth.font.bold = True
    
    p_campus = doc.add_paragraph()
    p_campus.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_campus.paragraph_format.line_spacing = 1.5
    p_campus.paragraph_format.space_before = Pt(24)
    p_campus.paragraph_format.space_after = Pt(0)
    
    r_camp = p_campus.add_run("Program Studi S1 Sistem Informasi\nFakultas Teknik dan Informatika\nUniversitas Bina Sarana Informatika\nSurakarta\n2026")
    r_camp.font.name = 'Times New Roman'
    r_camp.font.size = Pt(12)
    r_camp.font.bold = True
    
    doc.add_page_break()

    # ==================== 2. LEMBAR PERSETUJUAN ====================
    p_acc_head = doc.add_paragraph()
    p_acc_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_acc_head.paragraph_format.line_spacing = 1.5
    p_acc_head.paragraph_format.space_after = Pt(18)
    r_acc = p_acc_head.add_run("PERSETUJUAN PROPOSAL\nPRAKTIK KERJA LAPANGAN (PKL) / MAGANG")
    r_acc.font.name = 'Times New Roman'
    r_acc.font.size = Pt(14)
    r_acc.font.bold = True
    
    add_body_p("Praktik Kerja Lapangan (PKL) / Magang ini telah disetujui untuk dinilai pada Periode Semester 5 Tahun Akademik 2026/2027 di Program Studi S1 Sistem Informasi Fakultas Teknik dan Informatika Universitas Bina Sarana Informatika Surakarta.", indent=True)
    
    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_after = Pt(16)
    
    # Signature Table for Approval
    tbl_acc = doc.add_table(rows=3, cols=2)
    tbl_acc.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(tbl_acc)
    
    c00 = tbl_acc.cell(0, 0)
    c01 = tbl_acc.cell(0, 1)
    c10 = tbl_acc.cell(1, 0)
    c11 = tbl_acc.cell(1, 1)
    c20 = tbl_acc.cell(2, 0)
    c21 = tbl_acc.cell(2, 1)
    
    p = c00.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Menyetujui,\nDOSEN PENASEHAT AKADEMIK")
    r.font.name = 'Times New Roman'; r.font.size = Pt(12); r.font.bold = True
    
    p = c01.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Surakarta, 1 Agustus 2026\nPEMOHON / MAHASISWA")
    r.font.name = 'Times New Roman'; r.font.size = Pt(12); r.font.bold = True
    
    # Signature spaces
    p10 = c10.paragraphs[0]
    p10.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p10.paragraph_format.space_before = Pt(30)
    p10.paragraph_format.space_after = Pt(10)
    
    p11 = c11.paragraphs[0]
    p11.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p11.paragraph_format.space_before = Pt(4)
    p11.paragraph_format.space_after = Pt(4)
    if os.path.exists('ttd_fahrel.png'):
        r_ttd = p11.add_run()
        r_ttd.add_picture('ttd_fahrel.png', width=Inches(2.5))

    p = c20.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("( FRANSISCA NATALIA, S.E., M.M )\nNIP / NIDN: 200904405")
    r.font.name = 'Times New Roman'; r.font.size = Pt(12); r.font.bold = True
    
    p = c21.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("( MUHAMMAD FAHREL YULIYANTO )\nNIM. 19240513")
    r.font.name = 'Times New Roman'; r.font.size = Pt(12); r.font.bold = True
    
    doc.add_page_break()

    # ==================== 3. KATA PENGANTAR ====================
    p_kategori = doc.add_paragraph()
    p_kategori.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_kategori.paragraph_format.line_spacing = 1.5
    p_kategori.paragraph_format.space_after = Pt(12)
    r = p_kategori.add_run("KATA PENGANTAR")
    r.font.name = 'Times New Roman'; r.font.size = Pt(14); r.font.bold = True
    
    add_body_p("Puji dan syukur penulis panjatkan kehadirat Allah SWT atas segala rahmat, karunia, dan hidayah-Nya sehingga penulis dapat menyelesaikan proposal kegiatan Praktik Kerja Lapangan (PKL) / Magang ini dengan baik. Proposal ini disusun sebagai salah satu syarat akademis wajib semester 5 pada Program Studi S1 Sistem Informasi Fakultas Teknik dan Informatika Universitas Bina Sarana Informatika.")
    
    add_body_p("Proposal kegiatan magang ini diajukan untuk mendaftar pada posisi Junior Web Developer & IT Systems Intern di Unit IT / Sistem Informasi PT Kereta Api Indonesia (Persero) DAOP 6 Yogyakarta melalui portal InternHub KAI. Penulis berharap kegiatan magang ini dapat memberikan kontribusi riil bagi efisiensi operasional perusahaan sekaligus mengasah kompetensi teknis di dunia industri transportasi nasional.")
    
    add_body_p("Dalam penyusunan proposal ini, penulis menyampaikan rasa terima kasih dan penghargaan setinggi-tingginya kepada:")
    
    add_bullet_p("Jajaran Manajemen, Tim HRD, dan Kepala Unit IT PT Kereta Api Indonesia (Persero) DAOP 6 Yogyakarta atas kesempatan dan fasilitas pendaftaran magang yang diberikan.")
    add_bullet_p("Rektor Universitas Bina Sarana Informatika beserta Dekan Fakultas Teknik dan Informatika.")
    add_bullet_p("Ketua Program Studi S1 Sistem Informasi UBSI Kampus Surakarta beserta Ibu FRANSISCA NATALIA, S.E., M.M (200904405) selaku Dosen Penasehat Akademik yang senantiasa memberikan bimbingan dan arahan akademis.")
    add_bullet_p("Orang tua dan keluarga tercinta yang senantiasa memberikan doa, motivasi, dan dukungan moral maupun materiil.")
    add_bullet_p("Seluruh pihak yang telah memberikan bantuan dan dukungan dalam penyusunan proposal ini.")
    
    add_body_p("Penulis menyadari bahwa proposal kegiatan ini masih memiliki kekurangan. Oleh karena itu, kritik dan saran yang membangun sangat penulis harapkan demi kesempurnaan di masa mendatang. Semoga proposal ini dapat memberikan gambaran yang jelas mengenai rencana kegiatan magang dan memberikan manfaat bagi semua pihak.")
    
    # Signature block for Kata Pengantar
    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig.paragraph_format.line_spacing = 1.5
    p_sig.paragraph_format.space_before = Pt(12)
    p_sig.paragraph_format.space_after = Pt(0)
    
    r_s = p_sig.add_run("Surakarta, 1 Agustus 2026\nPenulis / Pemohon,\n\n")
    r_s.font.name = 'Times New Roman'; r_s.font.size = Pt(12); r_s.font.bold = True
    
    if os.path.exists('ttd_fahrel.png'):
        r_ttd2 = p_sig.add_run()
        r_ttd2.add_picture('ttd_fahrel.png', width=Inches(2.5))
        
    p_sig2 = doc.add_paragraph()
    p_sig2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig2.paragraph_format.line_spacing = 1.5
    p_sig2.paragraph_format.space_before = Pt(4)
    r_s2 = p_sig2.add_run("( MUHAMMAD FAHREL YULIYANTO )\nNIM : 19240513")
    r_s2.font.name = 'Times New Roman'; r_s2.font.size = Pt(12); r_s2.font.bold = True
    
    doc.add_page_break()

    # ==================== 4. DAFTAR ISI ====================
    p_di = doc.add_paragraph()
    p_di.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_di.paragraph_format.line_spacing = 1.5
    p_di.paragraph_format.space_after = Pt(12)
    r = p_di.add_run("DAFTAR ISI")
    r.font.name = 'Times New Roman'; r.font.size = Pt(14); r.font.bold = True
    
    add_toc_line("LEMBAR JUDUL PROPOSAL", "i", is_bold=True)
    add_toc_line("LEMBAR PERSETUJUAN PROPOSAL PKL", "ii", is_bold=True)
    add_toc_line("KATA PENGANTAR", "iii", is_bold=True)
    add_toc_line("DAFTAR ISI", "iv", is_bold=True)
    add_toc_line("BAB I PENDAHULUAN", "1", is_bold=True)
    add_toc_line("1.1 Latar Belakang", "1", is_bold=False, indent_cm=0.5)
    add_toc_line("1.2 Tujuan PKL / Magang", "2", is_bold=False, indent_cm=0.5)
    add_toc_line("1.3 Manfaat PKL / Magang", "2", is_bold=False, indent_cm=0.5)
    add_toc_line("BAB II TEMPAT DAN WAKTU PKL", "3", is_bold=True)
    add_toc_line("2.1 Tempat Kegiatan", "3", is_bold=False, indent_cm=0.5)
    add_toc_line("2.2 Waktu Pelaksanaan", "3", is_bold=False, indent_cm=0.5)
    add_toc_line("2.3 Bidang Pekerjaan & Rencana Program Kegiatan", "3", is_bold=False, indent_cm=0.5)
    add_toc_line("BAB III PENUTUP", "5", is_bold=True)
    add_toc_line("3.1 Kesimpulan & Harapan", "5", is_bold=False, indent_cm=0.5)
    add_toc_line("DAFTAR RIWAYAT HIDUP (CURRICULUM VITAE)", "6", is_bold=True)
    add_toc_line("LEMBAR PENILAIAN PRAKTIK KERJA LAPANGAN", "7", is_bold=True)

    doc.add_page_break()

    # ==================== BAB I PENDAHULUAN ====================
    add_bab_header("BAB I", "PENDAHULUAN")
    
    add_subbab_header("1.1 Latar Belakang")
    add_body_p("Perkembangan era digitalisasi dan teknologi informasi saat ini telah mengubah lanskap operasional berbagai industri nasional, tak terkecuali industri transportasi publik perkeretaapian. PT Kereta Api Indonesia (Persero) Daerah Operasi 6 Yogyakarta merupakan salah satu unit operasional vital yang melayani konektivitas transportasi massal di wilayah Daerah Istimewa Yogyakarta dan Jawa Tengah. Dalam mendukung efisiensi pelayanan, keandalan operasional kereta api, integrasi tiket elektronik, dan pemeliharaan infrastruktur, penerapan sistem informasi terintegrasi menjadi salah satu pilar utama strategi perusahaan.")
    
    add_body_p("Sebagai mahasiswa Semester 5 Program Studi S1 Sistem Informasi Fakultas Teknik dan Informatika Universitas Bina Sarana Informatika (UBSI) Kampus Surakarta, penulis memiliki minat dan fokus yang mendalam di bidang Fullstack Web Development (Laravel, PHP, JavaScript, Tailwind CSS, REST API), Mobile App Development (Flutter, Dart), Database Administration (SQL, PostgreSQL, MySQL), Software Quality Assurance (SQA Testing), serta integrasi Generative AI (Google Gemini, Google AI Studio, ChatGPT API). Keterampilan teknis ini juga telah diperkuat dengan perolehan Sertifikat BNSP Database Administrator (Reg: DMS.1241.00995 2026), Sertifikat Kompetensi SQA dari PT Tanjung Mulia Informatika, serta pengakuan 2 Hak Cipta (HKI) Resmi Kemenkumham RI dan perolehan Medal Silver (Silver Award) pada Pekan Inovasi Nasional Vol.3 2026 di Universitas Mercu Buana.")
    
    add_body_p("Untuk menjembatani keterkaitan dan kesepadanan (link and match) antara teori akademis di perguruan tinggi dengan kebutuhan nyata di dunia kerja industri BUMN, penulis mengajukan permohonan kegiatan Praktik Kerja Lapangan (PKL) / Magang pada posisi Junior Web Developer & IT Systems Intern di Unit IT / Sistem Informasi PT Kereta Api Indonesia (Persero) DAOP 6 Yogyakarta. Melalui kegiatan magang ini, penulis siap memberikan kontribusi positif dalam membantu pengembangan, pengujian, pemeliharaan, serta modernisasi sistem informasi operasional DAOP 6 Yogyakarta.")

    add_subbab_header("1.2 Tujuan PKL / Magang")
    add_body_p("Adapun tujuan dari pelaksanaan kegiatan Praktik Kerja Lapangan (PKL) / Magang ini adalah sebagai berikut:")
    add_bullet_p("Mengaplikasikan kompetensi teknis di bidang Web Development (Laravel, JavaScript), Mobile Development (Flutter, Dart), Database Administration, dan Generative AI pada ekosistem IT industri transportasi BUMN.")
    add_bullet_p("Membantu tim Unit IT DAOP 6 Yogyakarta dalam merancang, mengembangkan, menguji (SQA Testing), dan memelihara aplikasi web/sistem informasi internal perusahaan.")
    add_bullet_p("Memahami Standar Operasional Prosedur (SOP), tata kelola IT governance, dan siklus pengembangan perangkat lunak (SDLC) profesional di PT KAI (Persero).")
    add_bullet_p("Memenuhi syarat kurikulum akademis wajib Semester 5 pada Program Studi S1 Sistem Informasi Fakultas Teknik dan Informatika Universitas Bina Sarana Informatika.")

    add_subbab_header("1.3 Manfaat PKL / Magang")
    add_body_p("Pelaksanaan kegiatan magang ini diharapkan dapat memberikan manfaat yang saling menguntungkan bagi seluruh pihak terkait:")
    add_bullet_p("Mengembangkan hard skills dalam menangani sistem informasi skala besar (enterprise systems), melatih etos kerja, kedisiplinan, serta kemampuan memecahkan masalah (problem solving) di lingkungan kerja BUMN.", bold_prefix="a. Bagi Mahasiswa (Pemohon): ")
    add_bullet_p("Mendapatkan bantuan kontribusi langsung dari SDM mahasiswa bertalenta yang siap pakai di bidang Web/Mobile Development, Database, SQA, dan AI Integration untuk mempercepat penyelesaian tugas-tugas teknis operasional Unit IT.", bold_prefix="b. Bagi PT KAI DAOP 6 Yogyakarta: ")
    add_bullet_p("Mempererat sinergi dan kemitraan strategis antara perguruan tinggi dengan PT Kereta Api Indonesia (Persero) serta meningkatkan kualitas lulusan UBSI yang berstandar industri nasional.", bold_prefix="c. Bagi Perguruan Tinggi (UBSI): ")

    doc.add_page_break()

    # ==================== BAB II TEMPAT DAN WAKTU PKL ====================
    add_bab_header("BAB II", "TEMPAT DAN WAKTU PKL")
    
    add_subbab_header("2.1 Tempat Kegiatan")
    add_body_p("Kegiatan Praktik Kerja Lapangan (PKL) / Magang ini direncanakan bertempat pada:")
    
    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_info.paragraph_format.line_spacing = 1.5
    p_info.paragraph_format.left_indent = Inches(1.25 / 2.54)
    r = p_info.add_run(
        "Nama Perusahaan   : PT Kereta Api Indonesia (Persero) Daerah Operasi 6 Yogyakarta\n"
        "Unit Kerja Tujuan : Unit IT / Sistem Informasi (Information Technology)\n"
        "Posisi Dituju     : Junior Web Developer & IT Systems Intern\n"
        "Alamat Perusahaan : Jl. Lempuyangan No. 1, Bausasran, Danurejan, Kota Yogyakarta,\n"
        "                    Daerah Istimewa Yogyakarta 55212"
    )
    r.font.name = 'Times New Roman'; r.font.size = Pt(12)

    add_subbab_header("2.2 Waktu Pelaksanaan")
    add_body_p("Waktu pelaksanaan kegiatan PKL disesuaikan dengan kurikulum akademis kampus dan ketentuan pendaftaran InternHub KAI, yaitu berlangsung selama 3 (tiga) bulan:")
    
    p_waktu = doc.add_paragraph()
    p_waktu.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_waktu.paragraph_format.line_spacing = 1.5
    p_waktu.paragraph_format.left_indent = Inches(1.25 / 2.54)
    r = p_waktu.add_run(
        "Periode Magang   : 1 September 2026 s.d. 30 November 2026 (3 Bulan)\n"
        "Hari & Jam Kerja : Senin – Jumat (Mengikuti jam kerja operasional kantor PT KAI DAOP 6 Yogyakarta)"
    )
    r.font.name = 'Times New Roman'; r.font.size = Pt(12)

    add_subbab_header("2.3 Bidang Pekerjaan & Rencana Program Kegiatan")
    add_body_p("Sebagai Junior Web Developer & IT Systems Intern, rencana kegiatan magang disusun secara sistematis selama 12 minggu (3 bulan) sebagai berikut:")

    tbl_plan = doc.add_table(rows=1, cols=4)
    tbl_plan.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = tbl_plan.rows[0].cells
    hdr_titles = ["Bulan", "Minggu Ke-", "Rencana Program Kegiatan", "Output / Target Tahapan"]
    for i, title in enumerate(hdr_titles):
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.font.name = 'Times New Roman'; r.font.size = Pt(11); r.font.bold = True
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)

    plans_data = [
        ("Bulan I", "Minggu 1 - 2", "Orientasi lingkungan kerja Unit IT DAOP 6, pemahaman alur kerja operasional, pendataan kebutuhan sistem, dan analisis arsitektur basis data.", "Dokumen Analisis Kebutuhan Sistem & Laporan Orientasi"),
        ("Bulan I", "Minggu 3 - 4", "Rancangan antarmuka (UI/UX Figma) dan perancangan skema database relational (PostgreSQL/MySQL) untuk modul aplikasi web/mobile internal.", "Mockup Wireframe UI/UX & ERD Database"),
        ("Bulan II", "Minggu 5 - 6", "Pengembangan modul backend REST API (Laravel) dan frontend web/mobile (Flutter/Tailwind CSS) sesuai arahan supervisor pembimbing.", "Source Code Modul Aplikasi Web/Mobile Backend-Frontend"),
        ("Bulan II", "Minggu 7 - 8", "Integrasi sistem basis data, pengujian perangkat lunak (SQA Testing & Bug Fixing), dan uji keamanan akses data internal.", "Laporan Pengujian SQA & Berkas Sistem Siap UAT"),
        ("Bulan III", "Minggu 9 - 10", "Implementasi dan User Acceptance Testing (UAT) modul aplikasi bersama pengguna akhir di lingkungan kantor DAOP 6 Yogyakarta.", "Hasil UAT & Perbaikan Feedback Sistem"),
        ("Bulan III", "Minggu 11 - 12", "Penyusunan dokumentasi teknis (User Manual / API Documentation), penyerahan hasil karya ke Unit IT, dan penyelesaian laporan magang.", "User Manual Book & Laporan Akhir Magang PKL")
    ]

    for b, m, prog, outp in plans_data:
        row_cells = tbl_plan.add_row().cells
        data_tuple = (b, m, prog, outp)
        for i in range(4):
            p = row_cells[i].paragraphs[0]
            p.paragraph_format.line_spacing = 1.2
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if i in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(data_tuple[i])
            r.font.name = 'Times New Roman'; r.font.size = Pt(10.5)
            set_cell_margins(row_cells[i], top=100, bottom=100, left=120, right=120)

    doc.add_page_break()

    # ==================== BAB III PENUTUP ====================
    add_bab_header("BAB III", "PENUTUP")
    
    add_subbab_header("3.1 Kesimpulan & Harapan")
    add_body_p("Besar harapan penulis agar proposal kegiatan Praktik Kerja Lapangan (PKL) / Magang ini dapat diterima dan dipertimbangkan oleh Manajemen PT Kereta Api Indonesia (Persero) Daerah Operasi 6 Yogyakarta. Penulis berkomitmen penuh untuk bekerja secara profesional, mematuhi seluruh tata tertib dan Standar Operasional Prosedur (SOP) perusahaan, menjaga kerahasiaan data internal, serta berkontribusi nyata bagi Unit IT DAOP 6 Yogyakarta.")
    
    add_body_p("Demikian proposal ini dibuat dengan sebenar-benarnya sebagai syarat kelengkapan pendaftaran pada portal InternHub KAI. Atas perhatian, bimbingan, dan kesempatan yang diberikan oleh Manajemen PT Kereta Api Indonesia (Persero) DAOP 6 Yogyakarta, penulis mengucapkan terima kasih.")

    doc.add_page_break()

    # ==================== DAFTAR RIWAYAT HIDUP (CV) ====================
    p_cv_head = doc.add_paragraph()
    p_cv_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cv_head.paragraph_format.line_spacing = 1.5
    p_cv_head.paragraph_format.space_after = Pt(12)
    r = p_cv_head.add_run("DAFTAR RIWAYAT HIDUP (CURRICULUM VITAE)")
    r.font.name = 'Times New Roman'; r.font.size = Pt(14); r.font.bold = True
    
    add_subbab_header("A. BIODATA MAHASISWA")
    
    # Clean 2-column Table for Biodata to ensure PERFECT vertical colon alignment
    tbl_bio = doc.add_table(rows=11, cols=2)
    tbl_bio.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(tbl_bio)
    
    bio_fields = [
        ("NIM", ": 19240513"),
        ("Nama Lengkap", ": MUHAMMAD FAHREL YULIYANTO"),
        ("Tempat & Tanggal Lahir", ": Sragen, 22 Juli 2007"),
        ("Jenis Kelamin", ": Laki-laki"),
        ("Program Studi", ": S1 Sistem Informasi (Semester 5)"),
        ("Fakultas", ": Teknik dan Informatika"),
        ("Perguruan Tinggi", ": Universitas Bina Sarana Informatika (UBSI) Solo"),
        ("Alamat Asal", ": Krendowahono, Gondangrejo, Karanganyar, Jawa Tengah"),
        ("No. Telepon / WhatsApp", ": +62 889-8526-6525"),
        ("Email", ": muhmmadfahrel@gmail.com"),
        ("LinkedIn & Portofolio", ": linkedin.com/in/muhammad-fahrel | bit.ly/portofolio-fahrel")
    ]
    
    for row_idx, (label, val) in enumerate(bio_fields):
        row_cells = tbl_bio.rows[row_idx].cells
        c0 = row_cells[0]
        c1 = row_cells[1]
        
        c0.width = Inches(4.5 / 2.54)
        c1.width = Inches(9.5 / 2.54)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.line_spacing = 1.3
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(label)
        r0.font.name = 'Times New Roman'; r0.font.size = Pt(11); r0.font.bold = True
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.line_spacing = 1.3
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(val)
        r1.font.name = 'Times New Roman'; r1.font.size = Pt(11)
        
        set_cell_margins(c0, top=40, bottom=40, left=40, right=40)
        set_cell_margins(c1, top=40, bottom=40, left=40, right=40)

    add_subbab_header("B. RIWAYAT PENDIDIKAN FORMAL")
    add_bullet_p("Penerima Beasiswa BJU UBSI | IPK Sangat Memuaskan | Ketua Umum HIMSI UBSI Solo", bold_prefix="1. S1 Sistem Informasi — Universitas Bina Sarana Informatika (2024 – Sekarang / Smt 5): ")
    add_bullet_p("Lulusan Terbaik Multimedia | Juara 1 LKS Graphic Design Technology", bold_prefix="2. SMK Wihdatul Ummah — Jurusan Multimedia (2021 – 2024): ")

    add_subbab_header("C. PRESTASI & HAK CIPTA (HKI)")
    add_bullet_p("Medal Silver (Perak) Tingkat Nasional — Pekan Inovasi Nasional Vol.3 2026 Universitas Mercu Buana (Karya: Web Platform Terintegrasi Program Makan Bergizi Gratis MBG SPPG Nusukan).")
    add_bullet_p("Hak Cipta (HKI) Resmi Kemenkumham RI: Web System Terintegrasi Program Makan Bergizi Gratis (MBG) SPPG Nusukan (No: EC002026096415 - Juni 2026).")
    add_bullet_p("Hak Cipta (HKI) Resmi Kemenkumham RI: Kota Bersih Surakarta - Platform Smart Waste Management AI (No: EC002026053217 - April 2026).")
    add_bullet_p("Sertifikat Penghargaan Riset Perguruan Tinggi 2026 — BRIDA Kota Surakarta (Tim Fresh Tech).")
    add_bullet_p("Juara 2 Tingkat Nasional Video Reels Edukasi (SobiFest 2025) — Universitas BSI Jakarta.")
    add_bullet_p("Juara 1 LKS Graphic Design Technology 2024 — Tingkat Kabupaten Karanganyar.")

    add_subbab_header("D. SERTIFIKASI KOMPETENSI")
    add_bullet_p("Sertifikat BNSP — Database Administrator (LSP UBSI / BNSP No. Reg: DMS.1241.00995 2026).")
    add_bullet_p("Sertifikat Kompetensi Software Quality Assurance (SQA) — PT Tanjung Mulia Informatika (2026).")
    add_bullet_p("Pemakalah Seminar Nasional INOVATIF 2026 — UBSI & APTIKOM DKI Jakarta.")

    add_subbab_header("E. KEAHLIAN TEKNIS (SKILLS) & TOOLS")
    add_bullet_p("Flutter, Dart, PHP (Laravel), JavaScript, HTML/CSS, Tailwind CSS, REST API.", bold_prefix="Mobile & Web Dev : ")
    add_bullet_p("SQL, PostgreSQL, MySQL, SQA Manual Testing & Quality Assurance.", bold_prefix="Database & SQA   : ")
    add_bullet_p("Google Gemini, Google AI Studio, Google Flow, ChatGPT (OpenAI), Claude (Anthropic), Prompt Engineering.", bold_prefix="Generative AI    : ")
    add_bullet_p("Flutter SDK, Android Studio, Meta Ads Manager, CapCut, Adobe Illustrator, Figma, VS Code, GitHub.", bold_prefix="Tools & Platform : ")

    doc.add_page_break()

    # ==================== LEMBAR PENILAIAN ====================
    p_pen_head = doc.add_paragraph()
    p_pen_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pen_head.paragraph_format.line_spacing = 1.5
    p_pen_head.paragraph_format.space_after = Pt(12)
    r = p_pen_head.add_run("LEMBAR PENILAIAN PRAKTIK KERJA LAPANGAN (PKL)")
    r.font.name = 'Times New Roman'; r.font.size = Pt(14); r.font.bold = True
    
    add_body_p("Lembar penilaian ini diisi oleh Pembimbing Magang Instansi / Perusahaan PT Kereta Api Indonesia (Persero) DAOP 6 Yogyakarta pada akhir masa pelaksanaan magang:", indent=False)

    tbl_pen = doc.add_table(rows=1, cols=4)
    tbl_pen.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl_pen.rows[0].cells
    titles = ["No", "Unsur Penilaian", "Bobot (%)", "Nilai (0 - 100)"]
    for i, t in enumerate(titles):
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(t)
        r.font.name = 'Times New Roman'; r.font.size = Pt(11); r.font.bold = True
        set_cell_margins(hdr[i], top=100, bottom=100, left=100, right=100)

    scores_data = [
        ("1", "Kedisiplinan & Kehadiran (Punctuality)", "20%", ""),
        ("2", "Tanggung Jawab & Etos Kerja", "20%", ""),
        ("3", "Kemampuan Teknis (Web Dev / Database / SQA)", "25%", ""),
        ("4", "Kemampuan Memecahkan Masalah (Problem Solving)", "20%", ""),
        ("5", "Komunikasi & Kerjasama Tim (Teamwork)", "15%", ""),
        ("", "TOTAL NILAI AKHIR", "100%", "")
    ]

    for no, unsur, bobot, nilai in scores_data:
        row_cells = tbl_pen.add_row().cells
        vals = (no, unsur, bobot, nilai)
        for i in range(4):
            p = row_cells[i].paragraphs[0]
            p.paragraph_format.line_spacing = 1.2
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if i in [0, 2, 3]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(vals[i])
            r.font.name = 'Times New Roman'; r.font.size = Pt(10.5)
            if no == "":
                r.font.bold = True
            set_cell_margins(row_cells[i], top=100, bottom=100, left=100, right=100)

    add_body_p("Ketentuan Nilai: 80 - 100 (Sangat Baik / A) | 68 - 79 (Baik / B) | 56 - 67 (Cukup / C)", indent=False, space_after=12)

    # Signature block on Lembar Penilaian with current date
    p_eval_sig = doc.add_paragraph()
    p_eval_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_eval_sig.paragraph_format.line_spacing = 1.5
    p_eval_sig.paragraph_format.space_before = Pt(18)
    r = p_eval_sig.add_run(
        "Yogyakarta, 1 Agustus 2026\n"
        "Pembimbing Magang Instansi / Perusahaan\n"
        "PT Kereta Api Indonesia (Persero) DAOP 6 Yogyakarta"
    )
    r.font.name = 'Times New Roman'; r.font.size = Pt(12); r.font.bold = True

    # Save output document
    output_filename = 'Proposal_PKL_KAI_Daop6_Muhammad_Fahrel_Yuliyanto.docx'
    doc.save(output_filename)
    print(f"Proposal v2 successfully generated and saved to {output_filename}!")

if __name__ == '__main__':
    generate_ubsi_kai_proposal_v2()
