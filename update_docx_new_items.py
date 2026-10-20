import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def update_docx_new_items():
    doc = docx.Document('CV_Muhammad_Fahrel_Yuliyanto_2026.docx')
    table = doc.tables[0]
    c1 = table.rows[0].cells[1]

    # 1. Update PROFIL text
    for p in c1.paragraphs:
        if 'Mahasiswa Sistem Informasi semester 5' in p.text:
            p.text = (
                "Mahasiswa Sistem Informasi semester 5 di Universitas Bina Sarana Informatika (UBSI) "
                "dengan spesialisasi Mobile App & Web Development (Flutter, Dart, Laravel), Database Administration, "
                "Software Quality Assurance, Generative AI (Gemini, Google AI Studio, ChatGPT, Claude), dan Digital Marketing. "
                "Memiliki rekam jejak teruji dalam merancang ekosistem sistem informasi berbasis AI, Gamifikasi, dan Web, dibuktikan dengan pengakuan "
                "3 Hak Kekayaan Intelektual (HKI) resmi Kemenkumham (Platform MBG SPPG Surakarta, Smart Waste Management AI Kota Bersih Surakarta, "
                "& Web Official Desa Baran Klaten). Peraih Medal Silver (Silver Award) Pekan Inovasi Nasional Vol.3 2026 Universitas Mercu Buana, "
                "Diterima Magang di PT KAI (Persero) DAOP 6 Yogyakarta (Unit Sistem Informasi), dan Ketua Tim / Alumni BSI Explore 2026 (Lolos seleksi >2.000 peserta). "
                "Berpengalaman sebagai Pemateri Workshop AI/Content Creation, Narasumber Kampus, Ketua Umum HIMSI, serta praktisi Meta Ads dan "
                "kreator konten berprestasi tingkat nasional (Juara 2 SobiFest & Juara 1 LKS Desain Grafis)."
            )
            break

    doc.save('CV_Muhammad_Fahrel_Yuliyanto_2026.docx')
    print("PROFIL paragraph updated in DOCX with 3 HKIs and BSI Explore Team Leader!")

if __name__ == "__main__":
    update_docx_new_items()
